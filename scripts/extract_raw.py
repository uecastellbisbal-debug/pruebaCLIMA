#!/usr/bin/env python3
"""
Extract every sheet of the WM-KB Excel workbook into data/raw/*.json, unabridged.

Usage:
    pip install openpyxl
    python3 scripts/extract_raw.py /path/to/WMKB_WellnessMETER_Knowledge_Base.xlsx

Run this first whenever the source .xlsx changes, then run build_model.py.
"""
import argparse
import datetime
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT_RAW = ROOT / "data" / "raw"


def clean(v):
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def sheet_to_rows(ws):
    return [[clean(c) for c in row] for row in ws.iter_rows(values_only=True)]


def is_header_row(row):
    non_empty = [c for c in row if c not in (None, "")]
    if len(non_empty) < 2:
        return False
    return all(isinstance(c, str) for c in non_empty)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx_path", help="Path to the WM-KB .xlsx file")
    args = parser.parse_args()

    OUT_RAW.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(args.xlsx_path, data_only=True)

    manifest = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = sheet_to_rows(ws)

        while rows and all(c in (None, "") for c in rows[-1]):
            rows.pop()
        trim_to = 0
        for r in rows:
            for i in range(len(r) - 1, -1, -1):
                if r[i] not in (None, ""):
                    trim_to = max(trim_to, i + 1)
                    break
        rows = [r[:trim_to] for r in rows]

        records = None
        header = None
        if rows and is_header_row(rows[0]):
            raw_header = [h if h not in (None, "") else f"col_{i}" for i, h in enumerate(rows[0])]
            seen = {}
            header = []
            for h in raw_header:
                h = str(h)
                if h in seen:
                    seen[h] += 1
                    h = f"{h}_{seen[h]}"
                else:
                    seen[h] = 0
                header.append(h)
            records = []
            for r in rows[1:]:
                r = list(r) + [None] * (len(header) - len(r))
                rec = {header[i]: r[i] for i in range(len(header))}
                if any(v not in (None, "") for v in rec.values()):
                    records.append(rec)

        payload = {
            "sheet": name,
            "dimensions": ws.dimensions,
            "header": header,
            "rows_raw": rows if header is None else None,
            "records": records,
        }
        fname = re.sub(r"[^A-Za-z0-9_.\-]", "_", name)
        with open(OUT_RAW / f"{fname}.json", "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        manifest.append(
            {
                "sheet": name,
                "file": f"{fname}.json",
                "dimensions": ws.dimensions,
                "hasHeader": header is not None,
                "recordCount": len(records) if records is not None else len(rows),
            }
        )

    with open(OUT_RAW / "_manifest.json", "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)

    print(f"Extracted {len(wb.sheetnames)} sheets to {OUT_RAW}")


if __name__ == "__main__":
    main()
